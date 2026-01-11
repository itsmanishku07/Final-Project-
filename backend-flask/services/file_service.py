"""
File processing service for resume text extraction
Handles PDF and DOCX file processing using PyPDF2 and python-docx
"""

import logging
from typing import Optional
import PyPDF2
from docx import Document
import io

logger = logging.getLogger(__name__)

class FileProcessingService:
    """Service for processing uploaded resume files"""
    
    ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.doc'}
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    
    @staticmethod
    def is_allowed_file(filename: str, content_type: str) -> bool:
        """Check if file type is allowed"""
        if not filename:
            return False
        
        # Check file extension
        extension = filename.lower().split('.')[-1] if '.' in filename else ''
        if f'.{extension}' not in FileProcessingService.ALLOWED_EXTENSIONS:
            return False
        
        # Check content type
        allowed_content_types = {
            'application/pdf',
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'application/msword'
        }
        
        return content_type in allowed_content_types
    
    @staticmethod
    def is_valid_file_size(file_size: int) -> bool:
        """Check if file size is within limits"""
        return 0 < file_size <= FileProcessingService.MAX_FILE_SIZE
    
    @staticmethod
    def extract_text_from_file(file_content: bytes, content_type: str, filename: str) -> str:
        """Extract text from uploaded file"""
        try:
            if content_type == 'application/pdf':
                return FileProcessingService._extract_text_from_pdf(file_content)
            elif content_type in [
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'application/msword'
            ]:
                return FileProcessingService._extract_text_from_docx(file_content)
            else:
                raise ValueError(f"Unsupported file type: {content_type}")
                
        except Exception as e:
            logger.error(f"Failed to extract text from {filename}: {e}")
            raise RuntimeError(f"Text extraction failed: {e}")
    
    @staticmethod
    def _extract_text_from_pdf(file_content: bytes) -> str:
        """Extract text from PDF file"""
        try:
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text_content = []
            for page in pdf_reader.pages:
                text_content.append(page.extract_text())
            
            extracted_text = '\n'.join(text_content)
            
            if not extracted_text.strip():
                raise ValueError("No text content found in PDF")
            
            # Clean and normalize text
            cleaned_text = ' '.join(extracted_text.split())
            
            logger.debug(f"PDF text extracted successfully, length: {len(cleaned_text)}")
            return cleaned_text
            
        except Exception as e:
            logger.error(f"PDF text extraction failed: {e}")
            raise RuntimeError(f"PDF processing failed: {e}")
    
    @staticmethod
    def _extract_text_from_docx(file_content: bytes) -> str:
        """Extract text from DOCX file"""
        try:
            docx_file = io.BytesIO(file_content)
            doc = Document(docx_file)
            
            text_content = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text.strip())
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_content.append(cell.text.strip())
            
            extracted_text = '\n'.join(text_content)
            
            if not extracted_text.strip():
                raise ValueError("No text content found in DOCX")
            
            # Clean and normalize text
            cleaned_text = ' '.join(extracted_text.split())
            
            logger.debug(f"DOCX text extracted successfully, length: {len(cleaned_text)}")
            return cleaned_text
            
        except Exception as e:
            logger.error(f"DOCX text extraction failed: {e}")
            raise RuntimeError(f"DOCX processing failed: {e}")